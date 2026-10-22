"""Create the concise task DOCX using the bundled document Python runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

out = Path(__file__).resolve().parent
doc = Document()
for style in doc.styles:
    if style.element.pPr is not None:
        for border in list(style.element.pPr.findall(qn('w:pBdr'))):
            style.element.pPr.remove(border)
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.8)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1'):
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(11 if name == 'Normal' else 18 if name == 'Title' else 12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15
doc.styles['Subtitle'].font.italic = False
doc.add_paragraph('优化控制本轮任务说明', 'Title')
doc.add_paragraph('2026年10月22日', 'Subtitle')
doc.add_paragraph('本包交付新版十二输入模型的三机控制联调、随机来流训练、独立评估和中期答辩图表。全流程已跑通，MADDPG精度门槛尚未通过；当前结论限于代理模型内的软件验证。')
doc.add_paragraph('本轮完成内容', 'Heading 1')
doc.add_paragraph('接入队友新版PINN，统一十二维输入与六维输出；建立机舱朝向、桨距和转速执行器及P和PI对照。复用平稳湍流、慢变、阵风入口，先让MADDPG学习桨距，偏航跟踪与转速曲线由共同规则提供。')
doc.add_paragraph('三条独立训练流各完成16384步。验证选模后冻结策略，再开展新目标与新风轨迹测试。另提供7张Python绘制图及逐图讲解，适用于中期答辩。')
doc.add_paragraph('核心结果', 'Heading 1')
table = doc.add_table(rows=1, cols=2)
table.autofit = False
table.columns[0].width, table.columns[1].width = Inches(4.25), Inches(2.55)
table.rows[0].cells[0].text, table.rows[0].cells[1].text = '控制方法', '常规测试平均RMSE kW'
for name, value in [('初始MADDPG', '1083.62'), ('训练后MADDPG', '449.72 ± 40.11'), ('P控制', '552.37'), ('PI控制', '477.54')]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text = name, value
for i, row in enumerate(table.rows):
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcpr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for edge in ('top','left','bottom','right'):
            b=OxmlElement('w:'+edge); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
        tcpr.append(borders)
        margin=OxmlElement('w:tcMar')
        for edge in ('top','left','bottom','right'):
            b=OxmlElement('w:'+edge); b.set(qn('w:w'),'85'); b.set(qn('w:type'),'dxa'); margin.append(b)
        tcpr.append(margin)
        if i == 0:
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E8EEF3');tcpr.append(shade)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            for r in p.runs: r.bold = i == 0
doc.add_paragraph('误差为18场常规测试的场景RMSE均值；±为三条训练流均值的样本标准差。训练后平均误差较初始降低58.5%，但第0流、慢变和低风高目标仍弱于PI，不能宣称稳定优势。')
doc.add_paragraph('使用方式与后续重点', 'Heading 1')
doc.add_paragraph('先读readme.md，再看图表说明和评估报告。包内保留模型、策略、完整指标、来源清单和移动目录回放入口；历史优化工作由GitHub总索引及上轮归档追溯。')
doc.add_paragraph('下一步先诊断慢变跟踪和欠功率行为，再考虑开放转速与偏航学习。代理模型采用合成标签、执行器参数尚未标定，真实物理精度及发电收益仍需仿真或实测数据验证。')
doc.core_properties.title = '优化控制本轮任务说明'
doc.core_properties.author = '优化控制模块'
doc.save(out / '任务说明.docx')
print(out / '任务说明.docx')
