# -*- coding: utf-8 -*-
"""中期答辩 pptx 版面自检（无需 PowerPoint / LibreOffice）

做三件事，任何一条不过就退出码 1：

1. 越界：形状超出 16:9 画布（容差 0.01 in）
2. 溢出：按字号 + 字宽粗估文字所需高度，超过文本框高度 5% 即报
   —— 中文按 1.0 em，大写字母 0.60 em，数字 0.55 em，空格 0.28 em，其余 0.50 em
3. 重叠：两个非空文本框互相压住 >0.1 in × 0.1 in（表格单列，避免误报）

这是估算，不是排版真值：它能把「40pt 大字顶到下一行」这种事故提前抓出来，
但最终仍以 Keynote / PowerPoint 里的实际显示为准。

依赖：python-pptx
用法：python3 docs/defense/check_pptx_layout.py [docs/defense/中期答辩_v1_20261022.pptx]
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

EMU_PER_IN = 914400
TOL = 0.01        # 越界容差 / in
OVERFLOW = 1.05   # 高度容忍 5%
OVERLAP_MIN = 0.10


def char_em(ch: str) -> float:
    if ch.isspace():
        return 0.28
    if ord(ch) > 0x2E80:      # CJK / 全角标点
        return 1.0
    if ch.isupper():
        return 0.60
    if ch.isdigit():
        return 0.55
    return 0.50


def para_size(p) -> float:
    sizes = [r.font.size.pt for r in p.runs if r.font.size]
    if sizes:
        return max(sizes)
    return p.font.size.pt if p.font.size else 12.0


def needed_height(sh) -> float:
    """返回该文本框内容所需高度（pt）的估算值。"""
    width_pt = sh.width / EMU_PER_IN * 72.0
    avail = width_pt - 14.4          # 左右内边距各 0.1 in
    need = 0.0
    for p in sh.text_frame.paragraphs:
        size = para_size(p)
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.3
        lines, cur = 1, 0.0
        for ch in p.text:
            w = char_em(ch) * size
            if cur + w > avail and cur > 0:
                lines += 1
                cur = w
            else:
                cur += w
        need += lines * size * 1.2 * max(ls, 0.95)
        if p.space_before:
            need += p.space_before.pt
    return need


def check(path: Path) -> list[str]:
    prs = Presentation(str(path))
    W, H = prs.slide_width, prs.slide_height
    problems: list[str] = []
    for idx, slide in enumerate(prs.slides, 1):
        shapes = [s for s in slide.shapes if s.left is not None and s.width]
        for sh in shapes:
            right = sh.left + sh.width
            bottom = sh.top + sh.height
            if (sh.left < -TOL * EMU_PER_IN or sh.top < -TOL * EMU_PER_IN
                    or right > W + TOL * EMU_PER_IN or bottom > H + TOL * EMU_PER_IN):
                problems.append(f"p{idx:02d} 越界 {sh.shape_type} "
                                f"right={right / EMU_PER_IN:.2f}in bottom={bottom / EMU_PER_IN:.2f}in")
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            box_pt = sh.height / EMU_PER_IN * 72.0
            need = needed_height(sh)
            if need > box_pt * OVERFLOW:
                problems.append(f"p{idx:02d} 溢出 {sh.text_frame.text[:22]!r} "
                                f"需 ≈{need:.0f}pt / 框 {box_pt:.0f}pt")
        texts = [s for s in shapes if s.has_text_frame and s.text_frame.text.strip()]
        for i in range(len(texts)):
            for j in range(i + 1, len(texts)):
                a, b = texts[i], texts[j]
                ox = min(a.left + a.width, b.left + b.width) - max(a.left, b.left)
                oy = min(a.top + a.height, b.top + b.height) - max(a.top, b.top)
                if ox > OVERLAP_MIN * EMU_PER_IN and oy > OVERLAP_MIN * EMU_PER_IN:
                    problems.append(f"p{idx:02d} 重叠 {a.text_frame.text[:16]!r} × "
                                    f"{b.text_frame.text[:16]!r}")
    return problems


def main() -> int:
    here = Path(__file__).resolve().parent
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else here / "中期答辩_v1_20261022.pptx"
    if not target.exists():
        print(f"找不到 {target}；先运行 docs/defense/generate_midterm_pptx.py")
        return 2
    problems = check(target)
    prs = Presentation(str(target))
    print(f"{target.name}: {len(prs.slides._sldIdLst)} 页, {len(problems)} 个版面问题")
    for p in problems:
        print("  -", p)
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
