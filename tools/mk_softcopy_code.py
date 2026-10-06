#!/usr/bin/env python3
"""软著源代码 60 页文档生成器（前 30 页 + 后 30 页，每页 50 非空行）。

口径（中国版权保护中心，2026 现行）：
  - 源代码按「程序顺序」排定（tools/softcopy_code_order.txt 显式列出，可审计）；
  - 每页 50 行、不含空行（空行在排版前删除）；
  - 总代码 > 60 页时交前 30 + 后 30 页（第 30/31 页之间不连续）；
  - 第 60 页为模块结束页（= 程序最后一行）；
  - 页眉标注软件全称 + 版本号、右上角「第 X 页 共 60 页」；页脚标注权利人全称。

依赖：python3 + python-docx（pip3 install --user python-docx）
用法：
  python3 tools/mk_softcopy_code.py \
    --order tools/softcopy_code_order.txt \
    --header "多机协同智能功率控制与数字孪生演示平台 V1.0（待定）" \
    --footer "著作权人：（待定）" \
    --out docs/research/softcopy/output/源代码_60页_草稿.docx

产出：docx + 同名 .report.md（文件 → 页区间映射，用于核对 FLORIS 派生文件
不落入提交的 60 页；--derive 文件命中提交页时报 WARNING 并退出码 2）。
"""
from __future__ import annotations

import argparse
import math
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError:
    sys.exit("缺少 python-docx：pip3 install --user python-docx")

REPO = Path(__file__).resolve().parent.parent
LINES_PER_PAGE = 50
BLOCK_PAGES = 30


def read_order(order_path: Path) -> list[str]:
    files: list[str] = []
    for ln in order_path.read_text(encoding="utf-8").splitlines():
        ln = ln.strip()
        if not ln or ln.startswith("#"):
            continue
        files.append(ln)
    if not files:
        sys.exit(f"顺序表为空: {order_path}")
    return files


def src_files_under(src_dir: Path) -> set[str]:
    out = set()
    for p in src_dir.rglob("*"):
        if p.suffix in (".ts", ".tsx") and p.is_file():
            out.add(p.relative_to(REPO).as_posix())
    return out


def collect_lines(files: list[str]) -> tuple[list[str], dict[str, list[int]]]:
    """返回 (全部非空行, 文件 → 行号区间(1 基))。"""
    lines: list[str] = []
    spans: dict[str, list[int]] = {}
    for f in files:
        p = REPO / f
        if not p.is_file():
            sys.exit(f"顺序表文件不存在: {f}")
        first = len(lines) + 1
        for raw in p.read_text(encoding="utf-8").splitlines():
            if raw.strip():
                lines.append(raw.rstrip())
        spans[f] = [first, len(lines)]
    return lines, spans


def add_page_field(par) -> None:
    """向段落追加一个 Word PAGE 域（页码动态）。"""
    from docx.oxml import OxmlElement
    r1 = OxmlElement("w:r")
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    r1.append(f1)
    r2 = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = " PAGE "
    r2.append(it)
    r3 = OxmlElement("w:r")
    f3 = OxmlElement("w:fldChar")
    f3.set(qn("w:fldCharType"), "end")
    r3.append(f3)
    for el in (r1, r2, r3):
        par._p.append(el)


def build_docx(pages: list[list[str]], out: Path, header: str, footer: str) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.0)

    # 页眉：软件全称 + 版本号（左）… 制表位 … 第 [PAGE] 页 共 N 页（右）
    from docx.oxml import OxmlElement
    hpara = sec.header.paragraphs[0]
    hpara.text = ""
    r1 = hpara.add_run(header)
    r1.font.size = Pt(9)
    r1.font.name = "Times New Roman"
    r1._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    pPr = hpara._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), "9600")
    tabs.append(tab)
    pPr.append(tabs)
    hpara.add_run("\t")
    rp = hpara.add_run("第 ")
    rp.font.size = Pt(9)
    add_page_field(hpara)
    rt = hpara.add_run(f" 页 共 {len(pages)} 页")
    rt.font.size = Pt(9)

    # 页脚：权利人全称（居中）
    fpara = sec.footer.paragraphs[0]
    fpara.text = ""
    fpara.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf = fpara.add_run(footer)
    rf.font.size = Pt(9)
    rf.font.name = "Times New Roman"
    rf._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for i, page_lines in enumerate(pages, start=1):
        if i > 1:
            doc.add_page_break()
        for ln in page_lines:
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            r = p.add_run(ln)
            r.font.size = Pt(8.5)
            r.font.name = "Courier New"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--order", default="tools/softcopy_code_order.txt")
    ap.add_argument("--header", required=True, help="软件全称 + 版本号（页眉）")
    ap.add_argument("--footer", required=True, help="权利人全称（页脚）")
    ap.add_argument("--out", required=True)
    ap.add_argument(
        "--derive",
        default="twin/src/core/physics/gch.ts,twin/src/core/physics/florisTable.ts,twin/src/data/florisData.ts",
        help="FLORIS 派生文件（逗号分隔）；命中提交页则 WARNING",
    )
    args = ap.parse_args()

    order_path = REPO / args.order
    files = read_order(order_path)

    # 覆盖性：twin/src 下每个 ts/tsx 必须出现在顺序表（显式取舍，不许遗漏）
    missing = src_files_under(REPO / "twin" / "src") - set(files)
    if missing:
        sys.exit("以下 twin/src 文件未列入顺序表（必须显式处理）:\n  " + "\n  ".join(sorted(missing)))
    extra = [f for f in files if not f.startswith("twin/src/")]
    if any(f != "twin/scripts/selftest.mts" for f in extra):
        sys.exit(f"顺序表含非 twin/src 且非 selftest 的文件: {extra}")

    lines, spans = collect_lines(files)
    n = len(lines)
    total_pages = math.ceil(n / LINES_PER_PAGE)
    print(f"总非空行数: {n}  →  总页数: {total_pages}")
    if total_pages <= 60:
        print("总页数 ≤ 60：按规则应全部提交（本工具未实现该分支，请人工处理）")
        sys.exit(3)

    front = lines[: LINES_PER_PAGE * BLOCK_PAGES]
    back = lines[n - LINES_PER_PAGE * BLOCK_PAGES:]
    pages = [front[i * 50 : (i + 1) * 50] for i in range(BLOCK_PAGES)] + [
        back[i * 50 : (i + 1) * 50] for i in range(BLOCK_PAGES)
    ]

    # 页区间映射（程序页码：front → 1..30，back → 按程序行号折算）
    back_start = n - len(back)
    def doc_page_of(line_no: int) -> tuple[str, int | None]:
        if line_no <= len(front):
            return "提交·前30", math.ceil(line_no / 50)
        if line_no >= back_start + 1:
            return "提交·后30", 31 + (line_no - back_start - 1) // 50
        return "中间·不提交", None

    report = [
        f"# 软著源代码 60 页 —— 排版报告",
        "",
        f"- 总非空行数: {n}（总页数 {total_pages}，> 60，交前后各 30 页）",
        f"- 前 30 页 = 行 1–{len(front)}；后 30 页 = 行 {back_start+1}–{n}",
        f"- 第 30/31 页之间不连续（行 {len(front)} 与 {back_start+1} 之间省略 {back_start-len(front)} 行）",
        f"- 第 60 页末行 = 程序最后一行（模块结束页）✓",
        "",
        "| 文件 | 行区间 | 提交段 | 文档页 |",
        "|---|---|---|---|",
    ]
    derive = set(filter(None, args.derive.split(",")))
    warnings = []
    for f in files:
        lo, hi = spans[f]
        # 文件跨段时按端点分别标注
        segs = []
        for ln in (lo, hi):
            seg, dp = doc_page_of(ln)
            segs.append(f"{seg} p{dp}" if dp is not None else seg)
        seg_txt = " / ".join(sorted(dict.fromkeys(segs)))
        submitted = any(doc_page_of(ln)[0].startswith("提交") for ln in (lo, hi))
        if f in derive:
            if submitted:
                warnings.append(f"{f}（行 {lo}–{hi}）落入提交页！")
                row = f"| ⚠ {f} | {lo}–{hi} | **是（FLORIS 派生文件！）** | {seg_txt} |"
            else:
                row = f"| {f} | {lo}–{hi} | 否（中间段 ✓） | — |"
        elif submitted:
            row = f"| {f} | {lo}–{hi} | 是 | {seg_txt} |"
        else:
            row = f"| {f} | {lo}–{hi} | 否 | — |"
        report.append(row)

    out = Path(args.out)
    build_docx(pages, out, args.header, args.footer)
    report_path = out.with_suffix(".report.md")
    report.append("")
    report.append(f"派生文件核对: " + ("⚠ " + "; ".join(warnings) if warnings else "全部在中间不提交段 ✓"))
    report_path.write_text("\n".join(report) + "\n", encoding="utf-8")
    print(f"docx:  {out}")
    print(f"报告:  {report_path}")
    if warnings:
        print("WARNING: " + "; ".join(warnings), file=sys.stderr)
        sys.exit(2)
    print("派生文件全部在中间不提交段 ✓")


if __name__ == "__main__":
    main()
